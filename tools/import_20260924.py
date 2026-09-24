"""Import the 2026-09-24 C and review DOCX files into the deployed bank.

Uses researched answers for the eleven C-file conflicts, deduplicates complete
questions, and rejects structurally unrecoverable rows.
"""
from __future__ import annotations

import collections
import hashlib
import importlib.util
import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path

from docx import Document
from pypinyin import Style, lazy_pinyin

ROOT = Path(__file__).resolve().parents[1]
BANK_PATH = ROOT / 'docs/question-bank.json'
C_SOURCE = ROOT / 'data/交安C题库(2026-09-24).docx'
REVIEW_SOURCE = ROOT / 'data/交安复习题(2026-09-24).docx'
OLD_PREFIX = ('C20260924-', 'R20260924-')
TOC = re.compile(r'TOC\s*\\o\s*"1-5"\s*\\h\s*\\z\s*')
INLINE_ANSWER = re.compile(r'[（(]\s*(?:[A-H]{1,6}|正确|错误|对|错)\s*[)）]', re.I)
REVIEW_NOVEL = '根据《生产安全事故报告调查处理条例》规定，对于发生特别重大事故的单位，给予罚款的最高限额为'
VERIFIED_C_ANSWERS = {
    'DT00446': '正确',
    'DT00093': 'ACDE',
    'DT00067': 'BCDE',
    'DT00158': 'ABCE',
    'DT00124': 'BCD',
    'DT00418': 'AB',
    'DT00422': 'ABDE',
    'DT00355': 'CE',
    'C20260924-1280': 'AC',
    'DT00280': 'BCDE',
    'DT00282': 'ACDE',
}
VERIFIED_REVIEW_ANSWERS = {
    'DT00889-S04': 'B',
    'DT01103-Q': 'ABC',
}

# Recovered from the numbered source bank and checked against current law,
# official standards, or the complete equivalent already in the base bank.
# Keep the original C paragraph ordinal as the stable entry ID.
RECOVERED_C = {
    1: ('填空', '公路工程施工中，重大事故隐患排查治理的责任主体是( )。', [], '施工单位'),
    21: ('判断', '海因里希的因果连锁理论着重强调人的不安全行为和物的不安全状态在事故发生中的作用。( )', [], '正确'),
    31: ('判断', '破窗理论告诉我们，不安全行为往往受从众心理的影响。( )', [], '正确'),
    38: ('单选', '根据《建设工程安全生产管理条例》，总承包单位应当自行完成建设工程( )的施工。', ['整体结构', '主要结构', '所有结构', '主体结构'], 'D'),
    85: ('单选', '事故发生后，事故现场有关人员应当立即向本单位负责人报告；单位负责人接到报告后，应当于( )小时内向事故发生地县级以上人民政府有关部门报告。', ['1', '2', '12', '24'], 'A'),
    111: ('单选', '对因生产安全事故造成的职工死亡，一次性工亡补助金按全国上一年度城镇居民人均可支配收入的( )倍计算。', ['5', '10', '15', '20'], 'D'),
    177: ('单选', '生产经营单位主要负责人因生产安全事故受到刑事处罚或者撤职处分的，自刑罚执行完毕或者受处分之日起，( )年内不得担任任何生产经营单位的主要负责人。', ['1', '2', '5', '10'], 'C'),
    231: ('判断', '国务院应急管理部门和其他负有安全生产监督管理职责的部门应当根据各自的职责分工，制定相关行业、领域重大事故隐患的判定标准。( )', [], '正确'),
    238: ('判断', '与事故无关的单位或个人不用配合事故抢救。( )', [], '错误'),
    262: ('判断', '建设项目中的污染防治设施，应当与主体工程同时设计、同时施工、同时投产使用；设施应当符合经批准的生态环境影响报告书、生态环境影响报告表的要求，不得擅自拆除或者闲置。( )', [], '正确'),
    288: ('判断', '施工企业的从业人员均负有危险报告义务。( )', [], '正确'),
    318: ('判断', '建设行政主管部门或者其他有关部门不得将施工现场的监督检查委托给建设工程安全监督机构具体实施。( )', [], '错误'),
    346: ('判断', '用人单位应当依照法律、法规要求，严格遵守国家职业卫生标准，落实职业病预防措施，从源头上控制和消除职业病危害。( )', [], '正确'),
    426: ('单选', '生产经营单位的安全生产责任制大体可分为两个方面：一是( )方面各级人员的安全生产责任制；二是( )方面各职能部门的安全生产责任制。', ['生产；管理', '横向；纵向', '纵向；横向', '直接；间接'], 'C'),
    478: ('单选', '风险管理包括的最后一个过程是( )。', ['风险分析与评估过程', '风险控制对策的规划过程', '实施决策过程', '风险检查过程'], 'D'),
    501: ('判断', '从业人员发现事故隐患或者其他不安全因素，应当立即向现场人员或者本单位负责人报告。( )', [], '错误'),
    560: ('判断', '从防止事故的角度，安全隐患应当设法排除；尽早发现事故征兆并采取措施也有可能阻止事故发生，或者及时撤离以减少伤害和损失。( )', [], '正确'),
    567: ('判断', '某个可能发生的事件，其可能造成的损失程度和发生的概率都很大，则其风险量也越大。( )', [], '正确'),
    692: ('单选', '某桥墩高20m，上面架设的箱梁高4m，则该高处作业属于( )。', ['一级高处作业', '二级高处作业', '三级高处作业', '特级高处作业'], 'C'),
    693: ('单选', '高处作业人员上下应沿着( )行走。', ['立杆', '栏杆', '绳索', '扶梯'], 'D'),
    813: ('判断', '高处作业人员不得沿立杆或栏杆攀登；项目经理部应每天安排高处作业人员体检。( )', [], '错误'),
    824: ('判断', '作业人员在保管、加工、运输爆破器材过程中，严禁穿着化纤服装。( )', [], '正确'),
    934: ('判断', '蓄电池室、变压器室应有良好的通风。( )', [], '正确'),
}

# The complete versions are already present. Do not import a second copy of
# malformed Word rows merely to give them a new C-source ID.
RECOVERED_DUPLICATES = {
    117: 'DT00079', 491: 'C20260924-1433', 570: 'DT00564',
    591: 'DT00551', 598: 'DT00552', 605: 'DT00604', 606: 'DT00605',
    763: 'DT00586', 806: 'DT00494', 834: 'DT00401', 863: 'DT00661',
    866: 'DT00299', 1222: 'DT00202', 1226: 'DT00191',
}
RECOVERED_DUPLICATE_ANSWERS = {
    117: 'C', 491: '本质安全', 570: 'D', 591: 'A', 598: 'A',
    605: 'D', 606: 'A', 763: '错误', 806: '错误', 834: '正确',
    863: '错误', 866: '错误', 1222: 'BDE', 1226: 'ABCDE',
}


def clean(text: str) -> str:
    return re.sub(r'\s+', ' ', TOC.sub('', text)).strip()


def key(text: str) -> str:
    text = unicodedata.normalize('NFKC', clean(text))
    text = INLINE_ANSWER.sub('', text)
    return re.sub(r'[^\u4e00-\u9fa5a-z0-9]', '', text.lower())


def grams(text: str) -> set[str]:
    return {text[i:i+2] for i in range(len(text)-1)}


def full_text(stem: str, options: list[dict]) -> str:
    return key(stem + ' ' + ' '.join(option['text'] for option in options))


def entry_question(entry: dict) -> dict | None:
    qs = [b for b in entry['blocks'] if b['kind'] == 'question']
    return qs[0] if len(qs) == 1 else None


class Matcher:
    def __init__(self, entries):
        self.rows = []
        self.answers = {}
        self.questions = {}
        self.index = collections.defaultdict(set)
        for e in entries:
            q = entry_question(e)
            if q:
                self.add(e['id'], q)

    def add(self, id, q):
        self.answers[id] = q['answer']
        self.questions[id] = q
        stem = grams(key(q['stem']))
        whole = grams(full_text(q['stem'], q['options']))
        i = len(self.rows)
        self.rows.append((id, stem, whole))
        for gram in whole:
            self.index[gram].add(i)

    def nearest(self, stem, options):
        sg, fg = grams(key(stem)), grams(full_text(stem, options))
        overlaps = collections.Counter(i for g in fg for i in self.index[g])
        best = (0.0, 0.0, '')
        for i in sorted(overlaps):
            overlap = overlaps[i]
            id, other_sg, other_fg = self.rows[i]
            full_score = overlap / max(1, len(fg | other_fg))
            if full_score < best[0]:
                continue
            stem_score = len(sg & other_sg) / max(1, len(sg | other_sg))
            if (full_score, stem_score) > best[:2]:
                best = (full_score, stem_score, id)
        return best


def search_fields(q):
    raw = '\n'.join([q['stem'], *[o['text'] for o in q['options']], q['answer'], *q['notes']])
    normalize = lambda t: re.sub(r'\s+', '', unicodedata.normalize('NFKC', t)).lower()
    return {
        'text': normalize(raw),
        'pinyin': normalize(''.join(lazy_pinyin(raw, style=Style.NORMAL))),
        'initials': normalize(''.join(lazy_pinyin(raw, style=Style.FIRST_LETTER))),
    }


def refresh_entry_search(entry):
    parts=[]
    for block in entry['blocks']:
        if block['kind']=='question':
            parts.extend([block['stem'], *[o['text'] for o in block['options']], block['answer'], *block['notes']])
        else:
            parts.extend(block['paragraphs'])
    raw='\n'.join(parts)
    normalize=lambda t: re.sub(r'\s+', '', unicodedata.normalize('NFKC', t)).lower()
    entry['search']={'text':normalize(raw),
                     'pinyin':normalize(''.join(lazy_pinyin(raw,style=Style.NORMAL))),
                     'initials':normalize(''.join(lazy_pinyin(raw,style=Style.FIRST_LETTER)))}


def set_answer(entry, answer):
    q = entry_question(entry)
    assert q is not None
    old_answer = q['answer']
    q['answer'] = answer
    if old_answer != answer:
        q['stem'] = re.sub(r'([（(]\s*)' + re.escape(old_answer) + r'(\s*[)）])',
                           lambda match: match.group(1) + answer + match.group(2),
                           q['stem'], count=1)
    q['notes'] = []
    entry['title'] = q['stem'][:100]
    refresh_entry_search(entry)


def new_entry(id, number, kind, stem, options, answer, source, notes=None):
    q = {'kind':'question','id':id+'-Q','type':kind,'stem':clean(stem),
         'options':options,'answer':answer,'notes':notes or []}
    return {'id':id,'sourceNo':str(number),'sourceDocument':source,
            'type':kind,'section':source,'blocks':[q],
            'title':q['stem'][:100],'search':search_fields(q)}


def parse_c():
    rows = []
    for paragraph in Document(C_SOURCE).paragraphs[1:]:
        text = paragraph.text.strip()
        m = re.match(r'^(\d+)、(.*?)——\[([^]]+)\](?:\n|$)(.*)', text, re.S)
        if not m:
            rows.append({'error':'题目边界无法识别','raw':text})
            continue
        num, stem, kind, tail = m.groups()
        answer = re.search(r'正确答案[：:]\s*(.*)', tail)
        options = [{'key':k,'text':clean(v)} for k,v in re.findall(r'^\s*([A-H])\s+([^\n]+)', tail, re.M)]
        rows.append({'number':int(num),'stem':clean(stem),'kind':kind,
                     'options':options,'answer':clean(answer.group(1)) if answer else '', 'raw':text})
    return rows


def c_question(row):
    """Return a complete question or a concrete reason to hold the source row."""
    if 'error' in row:
        return None, row['error']
    stem, options, answer, kind = row['stem'], row['options'], row['answer'], row['kind']
    if not answer or not stem:
        return None, '题干或答案缺失'
    if row['number'] in (38, 426, 491):
        return None, '题干在源文档中截断'
    if kind == '填空题':
        if row['number'] in (1448, 1449):
            if answer not in ('A','B'):
                return None, '判断答案无法对应'
            return ('判断',re.sub(r'A[.．]正确B[.．]错误$', '', stem),[],
                    '正确' if answer=='A' else '错误'), None
        # Fill-in rows supply the answer text, but no option list.
        value = re.sub(r'\b[A-H][.．]\s*', '', answer).replace('|','；')
        if not value:
            return None, '填空答案只有选项字母'
        return ('填空',stem,[],value), None
    keys = [o['key'] for o in options]
    if len(keys) < 2 or len(set(keys)) != len(keys) or keys != [chr(65+i) for i in range(len(keys))]:
        return None, '选项残缺或标签错乱'
    if any(not o['text'] or re.search(r'\\o\b|TOC',o['text']) for o in options):
        return None, '选项文字残缺'
    if not re.fullmatch(r'[A-H]+', answer) or not set(answer).issubset(keys):
        return None, '答案不在选项内'
    if kind == '单选题' and len(answer) != 1:
        return None, '单选答案格式异常'
    if kind == '多选题' and len(answer) < 2:
        return None, '多选答案格式异常'
    if len(options)==2:
        labels = [key(x['text']) for x in options]
        if labels[0] in ('正','正确','对') and labels[1] in ('错','错误','误'):
            return ('判断',stem,[],'正确' if answer=='A' else '错误'), None
    return ('单选' if kind=='单选题' else '多选',stem,options,answer), None


def parse_review_novel():
    """The revised review DOCX contains one genuinely new question, twice."""
    spec = importlib.util.spec_from_file_location('legacy_parser', ROOT/'tools/parse_questions.py')
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    scratch = ROOT/'tmp'
    scratch.mkdir(exist_ok=True)
    paragraphs = [{'index':i,'text':p.text.strip()} for i,p in enumerate(Document(REVIEW_SOURCE).paragraphs) if p.text.strip()]
    mod.RAW = scratch/'review_20260924_raw.json'
    mod.OUT_QUESTIONS = scratch/'review_20260924_parsed.json'
    mod.OUT_REPORT = scratch/'review_20260924_report.json'
    mod.RAW.write_text(json.dumps(paragraphs,ensure_ascii=False),encoding='utf-8')
    mod.parse()
    parsed=json.loads(mod.OUT_QUESTIONS.read_text(encoding='utf-8'))['questions']
    found=[q for q in parsed if REVIEW_NOVEL in q.get('question','')]
    assert len(found)==2, f'Expected duplicate review question twice, found {len(found)}'
    assert all(q['answerText']=='D' and [o['key'] for o in q['options']]==list('ABCD') for q in found)
    return found[0],parsed


def main():
    bank=json.loads(BANK_PATH.read_text(encoding='utf-8'))
    entries=[e for e in bank['entries'] if not e['id'].startswith(OLD_PREFIX)]
    assert len(entries)==1206, 'The reviewed base bank has changed; reassess before importing'
    for entry in entries:
        changed=False
        for block in entry['blocks']:
            if block['kind']=='question':
                kept=[n for n in block['notes'] if not n.startswith(('【新版答案冲突】','【新版复习题答案冲突】'))]
                changed |= len(kept)!=len(block['notes'])
                block['notes']=kept
                if block['id'] in VERIFIED_REVIEW_ANSWERS:
                    assert block['answer'] == VERIFIED_REVIEW_ANSWERS[block['id']]
                    block['notes'] = []
                    changed = True
        if changed:
            refresh_entry_search(entry)
    by_id={e['id']:e for e in entries}
    # Some reviewed base questions are repeated under different IDs. Keep each
    # exact equivalent in sync so a search never returns the superseded answer.
    base_targets = {id: entry_question(by_id[id]) for id in VERIFIED_C_ANSWERS if id in by_id}
    for id, target in base_targets.items():
        signature = full_text(target['stem'], target['options'])
        for entry in entries:
            q = entry_question(entry)
            if q and full_text(q['stem'], q['options']) == signature:
                set_answer(entry, VERIFIED_C_ANSWERS[id])
    matcher=Matcher(entries)
    counts=collections.Counter()
    held=[]
    conflicts=[]
    for ordinal,row in enumerate(parse_c(),1):
        if ordinal in RECOVERED_DUPLICATES:
            assert row['number']==ordinal, f'C source paragraph moved: {ordinal}'
            counts['recovered_duplicate']+=1
            counts['duplicate']+=1
            continue
        if ordinal in RECOVERED_C:
            assert row['number']==ordinal, f'C source paragraph moved: {ordinal}'
            kind,stem,words,answer=RECOVERED_C[ordinal]
            options=[{'key':chr(65+i),'text':word} for i,word in enumerate(words)]
            candidate=(kind,stem,options,answer)
            reason=None
        else:
            candidate,reason=c_question(row)
        if reason:
            counts['held']+=1
            held.append({'sourceNo':row.get('number'),'row':ordinal,'reason':reason})
            continue
        kind,stem,options,answer=candidate
        score,stem_score,matched_id=matcher.nearest(stem,options)
        if score>=.90 and stem_score>=.80:
            counts['duplicate']+=1
            previous=matcher.answers[matched_id]
            canonical=lambda value: ''.join(sorted(re.match(r'^[A-H]+',value).group())) if re.match(r'^[A-H]+',value) else value
            matched=matcher.questions[matched_id]
            equivalent=False
            if not options and matched['options'] and re.fullmatch(r'[A-H]+',previous):
                existing_words='；'.join(o['text'] for o in matched['options'] if o['key'] in previous)
                equivalent=key(existing_words)==key(answer)
            if canonical(answer)!=canonical(previous) and not equivalent:
                conflicts.append({'sourceNo':row['number'],'row':ordinal,'sourceAnswer':answer,
                                  'existingId':matched_id,'existingAnswer':previous})
                counts['duplicate_answer_conflicts']+=1
            continue
        id=f'C20260924-{ordinal:04d}'
        notes=[]
        if row['number']==1435:
            notes=['【校订说明】源文档将《公路工程施工安全技术规范》归为“部门规章”，此分类与规范的行业标准性质不符；原答案仅供辨认原题。']
        e=new_entry(id,row['number'],kind,stem,options,answer,C_SOURCE.name,notes)
        if id in VERIFIED_C_ANSWERS:
            set_answer(e, VERIFIED_C_ANSWERS[id])
        entries.append(e)
        by_id[id]=e
        matcher.add(id,e['blocks'][0])
        counts['C_imported']+=1
        counts[f'C_{kind}']+=1
        if ordinal in RECOVERED_C:
            counts['recovered_imported']+=1
    review,review_parsed=parse_review_novel()
    review_conflicts={
        'DT00093-Q':'ACE',
        'DT00124-Q':'BCDE',
        'DT00280-Q':'ABCDE',
        'DT00889-S04':'A',
        'DT01103-Q':'ABCD',
    }
    question_index={q['id']:(entry,q) for entry in entries for q in entry['blocks'] if q['kind']=='question'}
    for qid,source_answer in review_conflicts.items():
        entry,oldq=question_index[qid]
        oldgrams=grams(key(oldq['stem']))
        sources=[q for q in review_parsed if ''.join(q.get('answer',[]))==source_answer and
                 len(oldgrams & grams(key(q.get('question',''))))/max(1,len(oldgrams | grams(key(q.get('question','')))))>=.85]
        assert sources, f'Missing reviewed answer conflict for {qid}'
        if entry['id'] not in VERIFIED_C_ANSWERS and qid not in VERIFIED_REVIEW_ANSWERS:
            note=(f'【新版复习题答案冲突】2026-09-24复习题答案为{source_answer}，'
                  f'原题库答案为{oldq["answer"]}；两份材料不一致，尚未统一。')
            oldq['notes'].append(note)
            refresh_entry_search(entry)
        counts['review_answer_conflicts']+=1
    note=('【校订说明】原题答案 D=2000 万元保留。题干所引《生产安全事故报告和调查处理条例》第37条为 500 万元上限；'
          '《安全生产法》第114条对负有责任单位规定 1000 万至 2000 万元，情节特别严重、影响特别恶劣时可按该数额的 2 至 5 倍罚款。')
    stem=clean(review['question'])
    options=[{'key':o['key'],'text':clean(o['text'])} for o in review['options']]
    score,stem_score,_=matcher.nearest(stem,options)
    assert score<.90 or stem_score<.80, 'Review question is already in the bank'
    e=new_entry('R20260924-0001',review['sourceNo'],'单选',stem,options,'D',REVIEW_SOURCE.name,[note])
    entries.append(e)
    counts['review_imported']+=1
    final_by_id={entry['id']:entry for entry in entries}
    for number,id in RECOVERED_DUPLICATES.items():
        assert any(block['kind']=='question' and block['answer']==RECOVERED_DUPLICATE_ANSWERS[number]
                   for block in final_by_id[id]['blocks']), (number,id)
    assert counts['recovered_imported']==len(RECOVERED_C)
    assert counts['recovered_duplicate']==len(RECOVERED_DUPLICATES)
    bank['entries']=entries
    bank['generatedDate']=datetime.now().astimezone().date().isoformat()
    bank['sourceDocuments']=[
        {'name':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
        for p in (ROOT/'交安复习题全部.大题层级修正版-20260907.docx',C_SOURCE,REVIEW_SOURCE)
    ]
    BANK_PATH.write_text(json.dumps(bank,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
    report={'counts':dict(counts),'totalEntries':len(entries),'held':held,'duplicateAnswerConflicts':conflicts}
    (ROOT/'tmp/import_20260924_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'counts':dict(counts),'totalEntries':len(entries),'held':held,'duplicateAnswerConflicts':conflicts},ensure_ascii=False))


if __name__=='__main__':
    main()
